"""
Vistas para la aplicación de Administrador.

Este módulo contiene todas las vistas relacionadas con las funcionalidades
disponibles para los usuarios con rol de administrador, incluyendo:
- Gestión de usuarios
- Seguimiento de pedidos
- Exportación de reportes (PDF y Excel)
"""

from django.shortcuts import render, get_object_or_404, redirect, reverse
from django.contrib.auth.decorators import login_required
from django.views.decorators.csrf import csrf_exempt
import json
from django.http import JsonResponse, HttpResponse
from core.models import Usuario, Rol, Solicitud, Borrador, CentroFormacion, Programa
from django.contrib import messages


@login_required
def panel_admin(request):
    """
    Vista del panel principal del administrador.
    
    Verifica que el usuario tenga rol de administrador antes de mostrar el panel.
    
    Args:
        request: Objeto HttpRequest del usuario autenticado
        
    Returns:
        HttpResponse: Renderiza el panel de administración o redirige si no tiene permisos
    """
    if request.user.rol is None or request.user.rol.nombre_rol != "administrador":
        return redirect(reverse("acceso_denegado"))
    return render(request, "administrador/panel_admin.html")


@login_required
def administracion_usuarios(request):
    """
    Vista para la administración de usuarios del sistema.
    
    Muestra la lista de todos los usuarios (excepto superusuario) y sus roles
    para que el administrador pueda gestionarlos.
    
    Args:
        request: Objeto HttpRequest del usuario autenticado
        
    Returns:
        HttpResponse: Renderiza la página de administración de usuarios
    """
    usuarios = Usuario.objects.filter(is_superuser=False)
    roles = Rol.objects.all()
    return render(request, "administrador/administracion_usuarios.html", {"usuarios": usuarios, "roles": roles})


@login_required
def historial_usuarios(request, id_usuario):
    """
    Vista para mostrar el historial de solicitudes de un usuario específico.
    
    Muestra todas las solicitudes realizadas por el usuario seleccionado,
    ordenadas por fecha de solicitud (más recientes primero).
    
    Args:
        request: Objeto HttpRequest del usuario autenticado
        id_usuario: ID del usuario cuyo historial se desea ver
        
    Returns:
        HttpResponse: Renderiza el historial de solicitudes del usuario
    """
    usuario = get_object_or_404(Usuario, id_usuario=id_usuario)
    solicitudes = usuario.solicitudes.all()
    
    # Ordenar por fecha de solicitud (más recientes primero)
    solicitudes = solicitudes.order_by("-fecha_solicitud")
    return render(
        request,
        "administrador/historial-de-usuario.html",
        {
            "usuario": usuario,
            "solicitudes": solicitudes
        }
    )




@login_required
def editar_usuario(request, id_usuario):
    """
    Vista para editar la información de un usuario.
    
    Permite al administrador modificar el nombre, apellido, correo y rol
    de un usuario. No permite editar usuarios inactivos.
    
    Args:
        request: Objeto HttpRequest con los datos del formulario
        id_usuario: ID del usuario a editar
        
    Returns:
        HttpResponseRedirect: Redirige a la administración de usuarios
    """
    usuario = get_object_or_404(Usuario, id_usuario=id_usuario)

    # Bloquear edición si está inactivo
    if not usuario.is_active:
        messages.error(request, "El usuario está inactivo y no puede ser editado.")
        return redirect("administracion-usuarios")

    if request.method == "POST":
        nombre = request.POST.get("nombre")
        apellido = request.POST.get("apellido")
        correo = request.POST.get("correo")
        rol_id = request.POST.get("rol")

        if nombre:
            usuario.nombre = nombre
        if apellido:
            usuario.apellido = apellido
        if correo:
            usuario.correo = correo
        if rol_id:
            usuario.rol_id = int(rol_id)

        usuario.save()
        messages.success(request, f"Usuario {usuario.nombre} {usuario.apellido} editado correctamente.")
        return redirect("administracion-usuarios")

    messages.warning(request, "Método no permitido.")
    return redirect("administracion-usuarios")


@login_required
def eliminar_usuario(request, id_usuario):
    """
    Vista para eliminar un usuario del sistema.
    
    Solo permite eliminar usuarios activos. Los usuarios inactivos
    no pueden ser eliminados.
    
    Args:
        request: Objeto HttpRequest (debe ser POST)
        id_usuario: ID del usuario a eliminar
        
    Returns:
        HttpResponseRedirect: Redirige a la administración de usuarios
    """
    if request.method == "POST":
        try:
            usuario = Usuario.objects.get(id_usuario=id_usuario)

            # Bloquear eliminación si está inactivo
            if not usuario.is_active:
                messages.warning(request, "No se puede eliminar un usuario inactivo.")
                return redirect("administracion-usuarios")

            usuario.delete()
            messages.success(request, "Usuario eliminado correctamente.")
        except Usuario.DoesNotExist:
            messages.error(request, "Usuario no encontrado.")
    else:
        messages.error(request, "Método no permitido.")
    
    return redirect("administracion-usuarios")


@login_required
def cambiar_estado_usuario(request, id_usuario):
    """
    Vista para activar o desactivar un usuario.
    
    Alterna el estado activo/inactivo del usuario. Los usuarios inactivos
    no pueden iniciar sesión en el sistema.
    
    Args:
        request: Objeto HttpRequest (debe ser POST)
        id_usuario: ID del usuario cuyo estado se desea cambiar
        
    Returns:
        HttpResponseRedirect: Redirige a la administración de usuarios
    """
    if request.method == "POST":
        try:
            usuario = Usuario.objects.get(id_usuario=id_usuario)

            # Cambiar estado activo/inactivo
            usuario.is_active = not usuario.is_active
            usuario.save()

            if usuario.is_active:
                messages.success(request, f"Usuario {usuario.nombre} {usuario.apellido} activado correctamente.")
            else:
                messages.info(request, f"Usuario {usuario.nombre} {usuario.apellido} desactivado correctamente.")

        except Usuario.DoesNotExist:
            messages.error(request, "Usuario no encontrado.")
    else:
        messages.error(request, "Método no permitido.")

    return redirect("administracion-usuarios")


from django.db.models import Q


@login_required
def seguimiento_pedidos(request):
    """
    Vista para el seguimiento de pedidos y solicitudes.
    
    Muestra todas las solicitudes del sistema con opciones de filtrado
    por búsqueda de texto y estado. También muestra los borradores guardados.
    
    Args:
        request: Objeto HttpRequest con parámetros de búsqueda opcionales
        
    Returns:
        HttpResponse: Renderiza la página de seguimiento de pedidos
    """
    buscar = request.GET.get("buscar", "").strip()
    estado = request.GET.get("estado", "").strip()

    solicitudes = Solicitud.objects.select_related("id_aprendiz", "id_producto").order_by("-fecha_solicitud")
    borradores = Borrador.objects.all().order_by("-creado")

    if buscar:
        solicitudes = solicitudes.filter(
            Q(id_aprendiz__nombre__icontains=buscar) |
            Q(id_producto__tipo__nombre__icontains=buscar)
        )

    if estado:
        solicitudes = solicitudes.filter(estado_solicitud__iexact=estado)

    return render(request, "administrador/seguimiento-de-pedidos.html", {
        "solicitudes": solicitudes,
        "borradores": borradores,
        "buscar": buscar,
        "estado": estado,
        "centros": CentroFormacion.objects.all()
    })







from django.http import HttpResponse
from django.db.models import Q
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, LongTable
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle

def exportar_pdf(request):
    """
    Vista para exportar solicitudes a formato PDF.
    
    Genera un reporte en PDF con las solicitudes filtradas según los parámetros
    de búsqueda, estado y usuario. El PDF se genera en formato landscape.
    
    Args:
        request: Objeto HttpRequest con parámetros de filtrado opcionales
        
    Returns:
        HttpResponse: Archivo PDF con el reporte de solicitudes
    """
    buscar = request.GET.get("buscar", "").strip().lower()
    estado = request.GET.get("estado", "").strip().lower()
    usuario = request.GET.get("usuario", "").strip()

    solicitudes = Solicitud.objects.select_related("id_aprendiz", "id_producto").all()

    if buscar:
        solicitudes = solicitudes.filter(
            Q(id_solicitud__icontains=buscar)
            | Q(fecha_solicitud__icontains=buscar)
            | Q(id_aprendiz__nombre__icontains=buscar)
            | Q(tipo_nombre__icontains=buscar)
            | Q(talla__icontains=buscar)
            | Q(color__icontains=buscar)
            | Q(cantidad__icontains=buscar)
            | Q(centro_formacion__icontains=buscar)
            | Q(programa__icontains=buscar)
            | Q(ficha__icontains=buscar)
            | Q(detalles_adicionales__icontains=buscar)
            | Q(fecha_finalizacion__icontains=buscar)
            | Q(estado_solicitud__icontains=buscar)
        )

    if estado:
        solicitudes = solicitudes.filter(estado_solicitud__iexact=estado)

    if usuario:
        solicitudes = solicitudes.filter(id_aprendiz__nombre__icontains=usuario)

    # 📌 Generar título dinámico
    if estado and usuario:
        titulo = f"Reporte de solicitudes {estado.title()} de {usuario.title()}"
    elif estado and buscar:
        titulo = f"Reporte de solicitudes {estado.title()} ({buscar})"
    elif usuario and buscar:
        titulo = f"Reporte de solicitudes de {usuario.title()} ({buscar})"
    elif estado:
        titulo = f"Reporte de solicitudes {estado.title()}"
    elif usuario:
        titulo = f"Reporte de solicitudes de {usuario.title()}"
    elif buscar:
        titulo = f"Reporte de solicitudes ({buscar})"
    else:
        titulo = "Reporte de todas las solicitudes"

    # 📄 Respuesta PDF
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = "attachment; filename=solicitudes.pdf"

    doc = SimpleDocTemplate(response, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elementos = []

    # 🏷️ Título
    elementos.append(Paragraph(titulo, styles["Title"]))
    elementos.append(Spacer(1, 12))

    # 📊 Datos de la tabla
    data = [[
        "ID", "Fecha", "Usuario", "Producto", "Talla", "Color", "Cantidad",
        "Centro", "Programa", "Ficha", "Detalles Adicionales" ,"Fecha Fin", "Estado"
    ]]

    estilo_celda = ParagraphStyle("Normal", fontSize=7, alignment=1)

    for s in solicitudes:
        data.append([
            str(s.id_solicitud),
             s.fecha_solicitud.strftime("%Y-%m-%d") if s.fecha_solicitud else "",
            str(s.id_aprendiz.get_full_name() if hasattr(s.id_aprendiz, "get_full_name") else s.id_aprendiz),
            str(s.tipo_nombre),
            str(s.talla_nombre),
            str(s.color_nombre),
            str(s.cantidad),
            str(s.centro_nombre),
            str(s.programa_nombre),
            str(s.ficha),
            Paragraph(str(s.detalles_adicionales), estilo_celda),
             s.fecha_finalizacion.strftime("%Y-%m-%d") if s.fecha_finalizacion else "",
            str(s.estado_solicitud),
        ])

    # 📐 Ajustar anchos de columnas
    colWidths = [
        35,   
        55,   
        105,  
        50,   
        30,   
        45,   
        40,   
        50,   
        50,   
        50,  
        100,  
        60,   
        60,   
    ]

    table = LongTable(data, repeatRows=1, colWidths=colWidths)

    # 🎨 Estilos de la tabla
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#c3f0ca")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
    ]))

    elementos.append(table)
    doc.build(elementos)

    return response



import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from django.http import HttpResponse
from django.db.models import Q
from django.utils.timezone import is_aware

def exportar_excel(request):
    """
    Vista para exportar solicitudes a formato Excel.
    
    Genera un archivo Excel (.xlsx) con las solicitudes filtradas según
    los parámetros de búsqueda, estado y usuario. Incluye formato de tabla.
    
    Args:
        request: Objeto HttpRequest con parámetros de filtrado opcionales
        
    Returns:
        HttpResponse: Archivo Excel con el reporte de solicitudes
    """
    buscar = request.GET.get("buscar", "").strip().lower()
    estado = request.GET.get("estado", "").strip().lower()
    usuario = request.GET.get("usuario", "").strip()

    solicitudes = Solicitud.objects.select_related("id_aprendiz", "id_producto").all()

    # 📌 Filtros
    if buscar:
        solicitudes = solicitudes.filter(
            Q(id_solicitud__icontains=buscar)
            | Q(fecha_solicitud__icontains=buscar)
            | Q(id_aprendiz__nombre__icontains=buscar)
            | Q(tipo_nombre__icontains=buscar)
            | Q(talla_nombre__icontains=buscar)
            | Q(color_nombre__icontains=buscar)
            | Q(cantidad__icontains=buscar)
            | Q(centro_nombre__icontains=buscar)
            | Q(programa_nombre__icontains=buscar)
            | Q(ficha__icontains=buscar)
            | Q(detalles_adicionales__icontains=buscar)
            | Q(fecha_finalizacion__icontains=buscar)
            | Q(estado_solicitud__icontains=buscar)
        )

    if estado:
        solicitudes = solicitudes.filter(estado_solicitud__iexact=estado)

    if usuario:
        solicitudes = solicitudes.filter(id_aprendiz__nombre__icontains=usuario)

    # 📌 Generar título dinámico
    if estado:
        titulo = f"Reporte de solicitudes {estado.title()}"
    elif usuario:
        titulo = f"Reporte de solicitudes de {usuario.title()}"
    elif buscar:
        titulo = f"Reporte de solicitudes (filtro {buscar})"
    else:
        titulo = "Reporte de todas las solicitudes"

    # Crear workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solicitudes"

    # Encabezados
    headers = [
        "ID Solicitud", "Fecha Solicitud", "Usuario", "Producto", "Talla", "Color", 
        "Cantidad", "Centro de Formación", "Programa", "Ficha",  "Detalles Adicionales",
        "Fecha Finalización", "Estado"
    ]
    ws.append(headers)

    # 📌 Agregar filas
    for solicitud in solicitudes:
        fecha_solicitud = solicitud.fecha_solicitud
        if is_aware(fecha_solicitud):
            fecha_solicitud = fecha_solicitud.replace(tzinfo=None)

        fecha_finalizacion = solicitud.fecha_finalizacion
        if fecha_finalizacion and is_aware(fecha_finalizacion):
            fecha_finalizacion = fecha_finalizacion.replace(tzinfo=None)

        ws.append([
            solicitud.id_solicitud,
            fecha_solicitud,
            solicitud.id_aprendiz.get_full_name() if solicitud.id_aprendiz else "",
            solicitud.tipo_nombre,
            solicitud.talla_nombre,
            solicitud.color_nombre,
            solicitud.cantidad,
            solicitud.centro_nombre,
            solicitud.programa_nombre,
            solicitud.ficha,
            solicitud.detalles_adicionales or "",
            fecha_finalizacion,
            solicitud.estado_solicitud.title(),
        ])

    # Tabla con estilo
    table = Table(displayName="TablaSolicitudes", ref=f"A1:M{ws.max_row}")
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=True
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    # 📌 Nombre dinámico del archivo
    filename = f"{titulo.replace(' ', '_').lower()}.xlsx"

    # Respuesta HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response









